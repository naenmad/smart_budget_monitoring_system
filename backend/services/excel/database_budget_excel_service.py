import io
import os
import re
from datetime import datetime, date
from decimal import Decimal
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from models.kategori import Kategori
from models.budget import Budget
from models.planning_header import PlanningHeader
from models.planning_detail import PlanningDetail
from models.pr_po_data import PrPoData
from models.preventive_cf import PreventiveCf
from models.upload_history import UploadHistory
from utils.db import db


class DatabaseBudgetExcelService:
    """
    Service untuk membaca (Import) dan menghasilkan (Export) file Excel
    'DATABASE WEB BUDGET.xlsx' yang terdiri dari 5 sheet:
      1. Summary
      2. Plan vs Actual (Kode)
      3. Plan vs Actual (Item)
      4. Out of Budget
      5. Preventive CF

    Kompatibel 100% dengan Microsoft Excel dan Apple Numbers.
    """

    FONT_NAME = "Calibri"

    # Styling Palettes (Clean, Modern & Professional)
    FONT_TITLE = Font(name=FONT_NAME, size=14, bold=True, color="1F4E78")
    FONT_SUBTITLE = Font(name=FONT_NAME, size=10, italic=True, color="595959")
    FONT_SECTION = Font(name=FONT_NAME, size=11, bold=True, color="1F4E78")
    FONT_HEADER = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
    FONT_DATA = Font(name=FONT_NAME, size=10, color="000000")
    FONT_DATA_BOLD = Font(name=FONT_NAME, size=10, bold=True, color="000000")
    FONT_TOTAL = Font(name=FONT_NAME, size=10, bold=True, color="1F4E78")

    FILL_HEADER_NAVY = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    FILL_HEADER_STEEL = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
    FILL_SUBTOTAL = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    FILL_ZEBRA = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")
    FILL_ALERT_SOFT = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")

    BORDER_THIN_GRAY = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9")
    )
    BORDER_TOTAL = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="1F4E78"),
        bottom=Side(style="double", color="1F4E78")
    )

    ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
    ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
    ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")
    ALIGN_HEADER = Alignment(horizontal="center", vertical="center", wrap_text=True)

    NUM_FORMAT_CURRENCY = "#,##0"
    NUM_FORMAT_CURRENCY_DEC = "#,##0.00"
    NUM_FORMAT_PERCENT = "0.0%"
    NUM_FORMAT_NUMBER = "#,##0"

    MONTHS_MAP = {
        1: "Jan-26", 2: "Feb-26", 3: "Mar-26", 4: "Apr-26",
        5: "Mei-26", 6: "Jun-26", 7: "Jul-26", 8: "Agu-26",
        9: "Sep-26", 10: "Okt-26", 11: "Nov-26", 12: "Des-26"
    }
    MONTH_NAMES_LIST = ["Jan-26", "Feb-26", "Mar-26", "Apr-26", "Mei-26", "Jun-26", "Jul-26", "Agu-26", "Sep-26", "Okt-26", "Nov-26", "Des-26"]

    @staticmethod
    def _clean_str(val):
        if val is None:
            return ""
        return str(val).strip()

    @staticmethod
    def _parse_num(val, default=0.0):
        if val is None or val == "":
            return default
        if isinstance(val, (int, float, Decimal)):
            return float(val)
        s = str(val).strip().replace("Rp", "").replace("IDR", "").replace(",", "").strip()
        try:
            return float(s)
        except ValueError:
            return default

    @staticmethod
    def _parse_int(val, default=None):
        if val is None or val == "":
            return default
        try:
            return int(float(val))
        except (ValueError, TypeError):
            return default

    @staticmethod
    def _find_sheet(wb, target_names):
        """
        Mencari sheet dengan pencocokan nama yang fleksibel
        (case-insensitive dan toleran terhadap spasi / variasi Apple Numbers).
        """
        for sheetname in wb.sheetnames:
            normalized = sheetname.strip().lower()
            for t in target_names:
                if normalized == t.strip().lower() or t.strip().lower() in normalized:
                    return wb[sheetname]
        return None

    # =========================================================================
    # IMPORT LOGIC
    # =========================================================================

    @classmethod
    def import_excel(cls, file_or_path, user_id=1, periode="2026"):
        """
        Membaca dan memproses workbook DATABASE WEB BUDGET.xlsx.
        Bisa membaca dari path file atau BytesIO/FileStorage.
        """
        if isinstance(file_or_path, str):
            wb = openpyxl.load_workbook(file_or_path, data_only=True)
            filename = os.path.basename(file_or_path)
        else:
            file_stream = io.BytesIO(file_or_path.read()) if hasattr(file_or_path, 'read') else file_or_path
            file_stream.seek(0)
            wb = openpyxl.load_workbook(file_stream, data_only=True)
            filename = getattr(file_or_path, 'filename', 'DATABASE_WEB_BUDGET.xlsx')

        # 1. Cari Sheet-Sheet Utama
        ws_summary = cls._find_sheet(wb, ["Summary", "Ringkasan"])
        ws_kode = cls._find_sheet(wb, ["Plan vs Actual (Kode)", "Plan vs Actual Kode", "Kode"])
        ws_item = cls._find_sheet(wb, ["Plan vs Actual (Item)", "Plan vs Actual Item", "Item"])
        ws_oob = cls._find_sheet(wb, ["Out of Budget", "Out of Plan", "OOB"])
        ws_prev = cls._find_sheet(wb, ["Preventive CF", "Preventive C/F", "Preventive"])

        if ws_item is None:
            raise ValueError("Sheet 'Plan vs Actual (Item)' tidak ditemukan. Pastikan format file sesuai struktur DATABASE WEB BUDGET.")

        # 2. Setup Kategori GL
        gl_categories = [
            ("50032902-123200", "SUPPLIES EXPENSES (QUALITY - QC)", "OPEX"),
            ("60290100-123200", "MARKETING EXPENSES (QUALITY - QC)", "OPEX"),
            ("60290500-123200", "TRANSPORTATION EXPENSE - LOCAL (QUALITY - QC)", "OPEX"),
            ("60290500-123000", "TRANSPORTATION EXPENSE - LOCAL (QUALITY - GNRL)", "OPEX"),
            ("50032802-123200", "REPAIR AND MAINTENANCE EXP OF PLANT EQUIPMENT", "OPEX"),
            ("50032902-123000", "SUPPLIES EXPENSES (QUALITY - GNRL)", "OPEX"),
            ("60290100-123000", "MARKETING EXPENSES (QUALITY - GNRL)", "OPEX"),
            ("60290300-123200", "SUPPLIES EXPENSES - OTHER (QUALITY - QC)", "OPEX"),
            ("60290600-123000", "TRANSPORTATION EXPENSE - OVERSEA (QUALITY - GNRL)", "OPEX"),
            ("60290600-123200", "TRANSPORTATION EXPENSE - OVERSEA (QUALITY - QC)", "OPEX"),
            ("60290800-123000", "STATIONARY G/A EXPENSE (QUALITY - GNRL)", "OPEX"),
            ("60290800-123200", "STATIONARY G/A EXPENSE (QUALITY - QC)", "OPEX"),
            ("60291100-123000", "UNIFORM EXPENSE (QUALITY - GNRL)", "OPEX"),
            ("E-6", "ENTERTAINMENT & RELATIONS", "OPEX"),
            ("OOP", "Out of Plan / Out of Budget", "OPEX"),
        ]

        kat_map = {}
        for kode, nama, tipe in gl_categories:
            kat = Kategori.query.filter_by(kode=kode).first()
            if not kat:
                kat = Kategori(kode=kode, nama=nama, tipe_formulir=tipe, is_active=True)
                db.session.add(kat)
                db.session.flush()
            else:
                kat.is_active = True
                kat.nama = nama
                kat.tipe_formulir = tipe
            kat_map[kode] = kat.id

        # 3. Create / Get Upload History & Planning Header
        unique_stored = f"{datetime.utcnow().strftime('%Y%m%d_%H%M%S')}_{filename}"
        upload_hist = UploadHistory(
            user_id=user_id,
            original_filename=filename,
            stored_filename=unique_stored,
            total_data=0,
            status="SUCCESS",
            uploaded_at=datetime.utcnow()
        )
        db.session.add(upload_hist)
        db.session.flush()

        header = PlanningHeader.query.filter_by(periode=periode).first()
        if not header:
            header = PlanningHeader(
                periode=periode,
                user_id=user_id,
                filename=filename,
                status="SUCCESS",
                uploaded_at=datetime.utcnow()
            )
            db.session.add(header)
            db.session.flush()
        else:
            header.filename = filename
            header.status = "SUCCESS"
            header.uploaded_at = datetime.utcnow()

        # 4. Parse Sheet 'Plan vs Actual (Item)'
        # Cari baris header dinamis: baris yang memiliki kolom 'no' dan 'pic' atau 'item planning'
        item_header_row = 4
        for r in range(1, 10):
            vals = [cls._clean_str(ws_item.cell(r, c).value).lower() for c in range(1, 10)]
            if any(v == "no" for v in vals) and any("pic" in v or "item planning" in v for v in vals):
                item_header_row = r
                break

        # Map header kolom
        col_map = {}
        for c in range(1, ws_item.max_column + 1):
            val = cls._clean_str(ws_item.cell(item_header_row, c).value).lower()
            if "no" == val: col_map["no"] = c
            elif "kode" in val: col_map["kode"] = c
            elif "account" in val: col_map["account"] = c
            elif "pic" in val: col_map["pic"] = c
            elif "item planning" in val or val == "item": col_map["item"] = c
            elif "actual" in val and "name" in val or "item name actual" in val: col_map["item_actual"] = c
            elif "bulan (no)" in val or "bulan no" in val: col_map["bulan_no"] = c
            elif "bulan rencana" in val or "bulan" in val and "no" not in val and "actual" not in val: col_map["bulan_rencana"] = c
            elif val == "qty" or "qty" in val and "actual" not in val: col_map["qty"] = c
            elif "actual qty" in val: col_map["actual_qty"] = c
            elif "nilai planning" in val or "planning" in val and "nilai" in val: col_map["nilai_plan"] = c
            elif "actual terealisasi" in val or "terealisasi" in val: col_map["act_real"] = c
            elif "bulan actual" in val: col_map["bulan_actual"] = c
            elif "selisih" in val: col_map["selisih"] = c
            elif "status" in val: col_map["status"] = c
            elif "keterangan" in val: col_map["keterangan"] = c
            elif "baris" in val: col_map["baris_ref"] = c
            elif "tipe" in val: col_map["tipe"] = c

        # Fallback jika kolom tidak terdeteksi
        c_no = col_map.get("no", 1)
        c_kode = col_map.get("kode", 2)
        c_acc = col_map.get("account", 3)
        c_pic = col_map.get("pic", 4)
        c_item = col_map.get("item", 5)
        c_act_item = col_map.get("item_actual", 6)
        c_bno = col_map.get("bulan_no", 7)
        c_brenc = col_map.get("bulan_rencana", 8)
        c_qty = col_map.get("qty", 9)
        c_act_qty = col_map.get("actual_qty", 10)
        c_plan = col_map.get("nilai_plan", 11)
        c_act_real = col_map.get("act_real", 12)
        c_bact = col_map.get("bulan_actual", 13)
        c_selisih = col_map.get("selisih", 14)
        c_stat = col_map.get("status", 15)
        c_ket = col_map.get("keterangan", 16)
        c_ref = col_map.get("baris_ref", 17)
        c_tipe = col_map.get("tipe", 18)

        # Putuskan link FK sementara pada pr_po_data agar tidak melanggar constraint
        PrPoData.query.filter(PrPoData.planning_detail_id.isnot(None)).update({"planning_detail_id": None})
        db.session.flush()

        # Hapus planning_detail lama untuk header ini
        PlanningDetail.query.filter_by(planning_header_id=header.id).delete()
        db.session.flush()

        plan_items_created = 0
        plan_detail_lookup = {}
        total_plan_by_kode = {}

        for r in range(item_header_row + 1, ws_item.max_row + 1):
            item_name = cls._clean_str(ws_item.cell(r, c_item).value)
            kode = cls._clean_str(ws_item.cell(r, c_kode).value)

            # Lewati baris total atau baris kosong
            if not item_name or item_name.upper() == "TOTAL" or not kode:
                continue

            pic = cls._clean_str(ws_item.cell(r, c_pic).value)
            act_item = cls._clean_str(ws_item.cell(r, c_act_item).value)
            b_rencana = cls._clean_str(ws_item.cell(r, c_brenc).value)
            qty = cls._parse_num(ws_item.cell(r, c_qty).value, default=None)
            act_qty = cls._parse_num(ws_item.cell(r, c_act_qty).value, default=None)
            nilai_plan = cls._parse_num(ws_item.cell(r, c_plan).value, default=0.0)
            act_real = cls._parse_num(ws_item.cell(r, c_act_real).value, default=0.0)
            b_actual = cls._clean_str(ws_item.cell(r, c_bact).value)
            selisih = cls._parse_num(ws_item.cell(r, c_selisih).value, default=act_real - nilai_plan)
            stat_str = cls._clean_str(ws_item.cell(r, c_stat).value)
            keterangan = cls._clean_str(ws_item.cell(r, c_ket).value)
            baris_ref = cls._clean_str(ws_item.cell(r, c_ref).value)
            tipe_val = cls._clean_str(ws_item.cell(r, c_tipe).value) or "Item"

            # Tentukan sub-kategori
            account = cls._clean_str(ws_item.cell(r, c_acc).value)
            if kode == "50032902-123200":
                if "PREVENTIVE" in baris_ref.upper() or "PREVENTIVE" in item_name.upper():
                    sub_kategori = "Preventive C/F"
                elif any(w in item_name.upper() for w in ["KALIBRASI", "CALIPER", "MICROMETER", "GAUGE", "THERMO"]):
                    sub_kategori = "Kalibrasi"
                else:
                    sub_kategori = "Consumables QC"
            elif "MARKETING" in account.upper():
                sub_kategori = "Marketing"
            elif "TRANSPORTATION" in account.upper():
                sub_kategori = "Transportation & BBM"
            else:
                sub_kategori = account

            # Status realisasi
            if "terealisasi" in stat_str.lower() and "belum" not in stat_str.lower():
                status_realisasi = "CLOSED"
            elif act_real > 0:
                status_realisasi = "PROSES"
            else:
                status_realisasi = "OPEN"

            remarks = f"PIC: {pic} | {keterangan}".strip(" | ")[:255]

            p_detail = PlanningDetail(
                planning_header_id=header.id,
                kategori_id=kat_map.get(kode, kat_map.get("OOP")),
                sub_kategori=sub_kategori,
                tipe=tipe_val,
                month=b_rencana,
                item=item_name,
                planning_amount=Decimal(str(nilai_plan)),
                remarks=remarks,
                status_realisasi=status_realisasi,
                pic=pic,
                item_actual=act_item,
                qty=Decimal(str(qty)) if qty is not None else None,
                actual_qty=Decimal(str(act_qty)) if act_qty is not None else None,
                actual_amount=Decimal(str(act_real)),
                bulan_actual=b_actual,
                selisih=Decimal(str(selisih)),
                status_detail=stat_str,
                baris_ref=baris_ref
            )
            db.session.add(p_detail)
            db.session.flush()

            plan_items_created += 1
            total_plan_by_kode[kode] = total_plan_by_kode.get(kode, 0.0) + nilai_plan

            # Store lookup for actual matching
            key_exact = (kode, item_name.lower(), b_rencana.lower())
            plan_detail_lookup[key_exact] = p_detail.id

        # 5. Update Master Budget
        # Jika sheet 'Plan vs Actual (Kode)' tersedia, ambil total planning tahunan per kode GL dari sana (13 kode resmi)
        budget_totals = {}
        if ws_kode:
            kode_h_row = 4
            for r in range(1, 10):
                vals = [cls._clean_str(ws_kode.cell(r, c).value).lower() for c in range(1, 6)]
                if any("kode" in v for v in vals) and any("planning" in v for v in vals):
                    kode_h_row = r
                    break
            for r in range(kode_h_row + 1, ws_kode.max_row + 1):
                kd = cls._clean_str(ws_kode.cell(r, 1).value)
                pl = cls._parse_num(ws_kode.cell(r, 6).value, 0.0)
                if kd and kd.upper() != "TOTAL":
                    budget_totals[kd] = budget_totals.get(kd, 0.0) + pl
        else:
            budget_totals = total_plan_by_kode

        for kode, tot_amt in budget_totals.items():
            kat_id = kat_map.get(kode)
            if kat_id:
                b_entry = Budget.query.filter_by(kategori_id=kat_id, periode=periode).first()
                if not b_entry:
                    b_entry = Budget(
                        kategori_id=kat_id,
                        periode=periode,
                        nominal=Decimal(str(tot_amt)),
                        created_by=user_id,
                        upload_id=upload_hist.id
                    )
                    db.session.add(b_entry)
                else:
                    b_entry.nominal = Decimal(str(tot_amt))
                    b_entry.upload_id = upload_hist.id

        # 6. Parse Sheet 'Out of Budget'
        oob_count = 0
        if ws_oob:
            oob_header_row = 3
            for r in range(1, 6):
                vals = [cls._clean_str(ws_oob.cell(r, c).value).lower() for c in range(1, 10)]
                if any("harga (gl)" in v or "harga" in v for v in vals):
                    oob_header_row = r
                    break

            # Bersihkan transaksi OOP lama agar selaras dengan file resmi
            PrPoData.query.filter(PrPoData.budget_status == "OOP").delete()
            db.session.flush()

            for r in range(oob_header_row + 1, ws_oob.max_row + 1):
                item_oob = cls._clean_str(ws_oob.cell(r, 7).value)
                if not item_oob or item_oob.upper() == "TOTAL":
                    continue

                bulan_str = cls._clean_str(ws_oob.cell(r, 2).value)
                kode_oob = cls._clean_str(ws_oob.cell(r, 3).value)
                sources = cls._clean_str(ws_oob.cell(r, 5).value) or "PR"
                tgl_gl = ws_oob.cell(r, 6).value
                harga_gl = cls._parse_num(ws_oob.cell(r, 8).value, 0.0)
                status_vs_get = cls._clean_str(ws_oob.cell(r, 11).value)
                pr_doc = cls._clean_str(ws_oob.cell(r, 12).value)
                part_no = cls._clean_str(ws_oob.cell(r, 13).value)
                item_get = cls._clean_str(ws_oob.cell(r, 14).value)
                po_get = cls._clean_str(ws_oob.cell(r, 15).value)
                gr_get = cls._clean_str(ws_oob.cell(r, 16).value)
                rcp_date = ws_oob.cell(r, 17).value
                nilai_get = cls._parse_num(ws_oob.cell(r, 18).value, None)
                selisih_get = cls._parse_num(ws_oob.cell(r, 19).value, None)
                catatan_comp = cls._clean_str(ws_oob.cell(r, 21).value)

                # Format request date
                req_date = tgl_gl.date() if isinstance(tgl_gl, datetime) else None
                if not req_date and isinstance(tgl_gl, date):
                    req_date = tgl_gl
                if not req_date:
                    # fallback ke tanggal 1 bulan terkait
                    m_num = 1
                    for idx, m_name in cls.MONTHS_MAP.items():
                        if m_name.lower() in bulan_str.lower():
                            m_num = idx
                            break
                    req_date = date(int(periode), m_num, 1)

                rcp_d = rcp_date.date() if isinstance(rcp_date, datetime) else (rcp_date if isinstance(rcp_date, date) else None)

                gl_acc = kode_oob.split("-")[0] if "-" in kode_oob else kode_oob
                cost_center = kode_oob.split("-")[1] if "-" in kode_oob else ""
                kat_id = kat_map.get(kode_oob, kat_map.get("OOP"))

                # Cek apakah record sudah ada berdasarkan doc & item
                doc_key = pr_doc or f"OOB-{bulan_str}-{r}"
                existing_pr = PrPoData.query.filter(
                    PrPoData.description == item_oob,
                    PrPoData.budget_status == "OOP",
                    PrPoData.total_price == Decimal(str(harga_gl))
                ).first()

                if not existing_pr:
                    pr_obj = PrPoData(
                        upload_id=upload_hist.id,
                        requisition_id=f"OOB-2026-{r:03d}",
                        pr_doc_num=doc_key,
                        po_doc_num=po_get or doc_key,
                        request_date=req_date,
                        order_date=req_date,
                        description=item_oob,
                        comment_text=catatan_comp,
                        supplier_name="Quality",
                        source=sources,
                        gl_account=gl_acc,
                        cost_center=cost_center,
                        qty=1,
                        uom="PCS",
                        unit_price=Decimal(str(harga_gl)),
                        total_price=Decimal(str(harga_gl)),
                        gr_legal_number=gr_get,
                        invoice=doc_key,
                        invoice_date=req_date,
                        pr_status="APPROVED",
                        po_status="CLOSED",
                        non_stock_item=1,
                        kategori_id=kat_id,
                        status_ai="DONE",
                        procurement_status="COMPLETED" if gr_get else "PO_ISSUED",
                        budget_status="OOP",
                        status_detail="OUT OF BUDGET",
                        beda_bulan=False,
                        doc_number=doc_key,
                        bulan=req_date.month,
                        part_number=part_no,
                        item_get_report=item_get,
                        po_get_report=po_get,
                        gr_get_report=gr_get,
                        receipt_date_get_report=rcp_d,
                        nilai_get_report=Decimal(str(nilai_get)) if nilai_get is not None else None,
                        selisih_get_report=Decimal(str(selisih_get)) if selisih_get is not None else None,
                        status_get_report=status_vs_get,
                        catatan_perbandingan=catatan_comp
                    )
                    db.session.add(pr_obj)
                    oob_count += 1
                else:
                    # Update fields
                    existing_pr.part_number = part_no
                    existing_pr.item_get_report = item_get
                    existing_pr.po_get_report = po_get
                    existing_pr.gr_get_report = gr_get
                    existing_pr.receipt_date_get_report = rcp_d
                    existing_pr.nilai_get_report = Decimal(str(nilai_get)) if nilai_get is not None else None
                    existing_pr.selisih_get_report = Decimal(str(selisih_get)) if selisih_get is not None else None
                    existing_pr.status_get_report = status_vs_get
                    existing_pr.catatan_perbandingan = catatan_comp

        # 7. Parse Sheet 'Preventive CF' (Tabel Rincian baris 20+)
        pcf_count = 0
        if ws_prev:
            PreventiveCf.query.delete()
            pcf_header_row = 19
            for r in range(1, 25):
                vals = [cls._clean_str(ws_prev.cell(r, c).value).lower() for c in range(1, 10)]
                if any("bulan list" in v for v in vals) or any("total list" in v for v in vals):
                    pcf_header_row = r
                    break

            for r in range(pcf_header_row + 1, ws_prev.max_row + 1):
                item_pcf = cls._clean_str(ws_prev.cell(r, 4).value)
                if not item_pcf or item_pcf.upper() == "TOTAL":
                    continue

                no_val = cls._parse_int(ws_prev.cell(r, 1).value)
                b_list = cls._clean_str(ws_prev.cell(r, 2).value)
                b_no = cls._parse_int(ws_prev.cell(r, 3).value)
                spek = cls._clean_str(ws_prev.cell(r, 5).value)
                qty = cls._parse_num(ws_prev.cell(r, 6).value, None)
                unit = cls._clean_str(ws_prev.cell(r, 7).value)
                cost = cls._parse_num(ws_prev.cell(r, 8).value, 0.0)
                tot_list = cls._parse_num(ws_prev.cell(r, 9).value, 0.0)
                remark_f = cls._clean_str(ws_prev.cell(r, 10).value)
                act_gl = cls._parse_num(ws_prev.cell(r, 11).value, 0.0)
                b_act = cls._clean_str(ws_prev.cell(r, 12).value)
                jml_tx = cls._parse_int(ws_prev.cell(r, 13).value, 0)
                selisih = cls._parse_num(ws_prev.cell(r, 14).value, 0.0)
                stat = cls._clean_str(ws_prev.cell(r, 15).value)
                ket = cls._clean_str(ws_prev.cell(r, 16).value)
                sumber = cls._clean_str(ws_prev.cell(r, 17).value)
                tgl_act = cls._clean_str(ws_prev.cell(r, 18).value)
                item_gl = cls._clean_str(ws_prev.cell(r, 19).value)
                doc_gl = cls._clean_str(ws_prev.cell(r, 20).value)

                pcf_obj = PreventiveCf(
                    header_id=header.id,
                    no_urut=no_val,
                    bulan_list=b_list,
                    bulan_no=b_no,
                    item=item_pcf,
                    spesifikasi=spek,
                    qty=Decimal(str(qty)) if qty is not None else None,
                    unit=unit,
                    cost=Decimal(str(cost)),
                    total_list=Decimal(str(tot_list)),
                    remark_file=remark_f,
                    actual_gl=Decimal(str(act_gl)),
                    bulan_actual=b_act,
                    jumlah_transaksi=jml_tx,
                    selisih=Decimal(str(selisih)),
                    status=stat,
                    keterangan=ket,
                    sumber=sumber,
                    tanggal_actual=tgl_act,
                    item_gl=item_gl,
                    no_dokumen_gl=doc_gl
                )
                db.session.add(pcf_obj)
                pcf_count += 1

        upload_hist.total_data = plan_items_created + oob_count + pcf_count
        db.session.commit()

        return {
            "success": True,
            "message": f"Sinkronisasi berhasil! {plan_items_created} item planning, {len(total_plan_by_kode)} master budget, {oob_count} Out of Budget, dan {pcf_count} Preventive CF berhasil dimuat.",
            "stats": {
                "planning_items": plan_items_created,
                "budget_categories": len(total_plan_by_kode),
                "out_of_budget_records": oob_count,
                "preventive_cf_records": pcf_count
            }
        }

    # =========================================================================
    # EXPORT LOGIC (Identik dengan DATABASE WEB BUDGET.xlsx, Full 5 Sheets)
    # =========================================================================

    @classmethod
    def export_excel(cls, periode="2026") -> io.BytesIO:
        """
        Menghasilkan file Excel 5-sheet yang persis identik dengan struktur
        DATABASE WEB BUDGET.xlsx, dengan styling profesional, kompatibilitas
        penuh terhadap Microsoft Excel dan Apple Numbers.
        """
        wb = openpyxl.Workbook()
        # openpyxl membuat default sheet 'Sheet', kita ganti namanya
        ws_summary = wb.active
        ws_summary.title = "Summary"

        ws_kode = wb.create_sheet(title="Plan vs Actual (Kode)")
        ws_item = wb.create_sheet(title="Plan vs Actual (Item)")
        ws_oob = wb.create_sheet(title="Out of Budget")
        ws_prev = wb.create_sheet(title="Preventive CF")

        # Aktifkan gridlines untuk semua sheets (Apple Numbers & Excel best practice)
        for ws in [ws_summary, ws_kode, ws_item, ws_oob, ws_prev]:
            ws.views.sheetView[0].showGridLines = True

        # Fetch Data dari Database
        header = PlanningHeader.query.filter_by(periode=periode).first()
        header_id = header.id if header else 0

        details = PlanningDetail.query.filter_by(planning_header_id=header_id).order_by(PlanningDetail.id.asc()).all()
        oob_list = PrPoData.query.filter(PrPoData.budget_status == "OOP").order_by(PrPoData.id.asc()).all()
        pcf_list = PreventiveCf.query.order_by(PreventiveCf.no_urut.asc()).all()
        kat_all = {k.id: k for k in Kategori.query.all()}

        # -------------------------------------------------------------
        # SHEET 3: Plan vs Actual (Item) (Ditulis duluan untuk lookup / referensi)
        # -------------------------------------------------------------
        ws_item.cell(row=1, column=1, value=f"Planning vs Actual per Item - QUALITY ({periode})").font = cls.FONT_TITLE
        ws_item.cell(row=2, column=1, value="Planning = ORIGINAL OPEX (Q1), per item per bulan rencana. Actual = transaksi GL yang cocok dengan item tersebut. Item planning level kode dievaluasi di sheet Plan vs Actual (Kode).").font = cls.FONT_SUBTITLE

        item_headers = [
            "No", "Kode (Trial Balance Report-Detail)", "Account", "PIC",
            "Item Planning", "Item Name Actual", "Bulan (no)", "Bulan Rencana",
            "Qty", "Actual Qty", "Nilai Planning", "Actual Terealisasi",
            "Bulan Actual", "Selisih (Actual - Planning)", "Status",
            "Keterangan", "Baris di ORIGINAL OPEX (Q1)", "Tipe"
        ]

        ws_item.row_dimensions[4].height = 28
        for col_idx, h in enumerate(item_headers, 1):
            cell = ws_item.cell(row=4, column=col_idx, value=h)
            cell.font = cls.FONT_HEADER
            cell.fill = cls.FILL_HEADER_NAVY
            cell.alignment = cls.ALIGN_HEADER
            cell.border = cls.BORDER_THIN_GRAY

        tot_plan_sum = 0.0
        tot_act_sum = 0.0
        curr_row = 5

        # Bulan no map
        month_no_map = {
            "jan": 1, "feb": 2, "mar": 3, "apr": 4, "may": 5, "mei": 5,
            "jun": 6, "jul": 7, "aug": 8, "agu": 8, "sep": 9, "oct": 10,
            "okt": 10, "nov": 11, "dec": 12, "des": 12
        }

        # Data rows
        for idx, d in enumerate(details, 1):
            k_obj = kat_all.get(d.kategori_id)
            kode = k_obj.kode if k_obj else ""
            acc = k_obj.nama.replace(" (QUALITY - QC)", "").replace(" (QUALITY - GNRL)", "") if k_obj else ""

            m_str = d.month or ""
            b_prefix = m_str[:3].lower()
            b_no = month_no_map.get(b_prefix, 1)

            p_amt = float(d.planning_amount or 0.0)
            a_amt = float(d.actual_amount or 0.0)
            diff = a_amt - p_amt

            tot_plan_sum += p_amt
            tot_act_sum += a_amt

            row_data = [
                idx,
                kode,
                acc,
                d.pic or "-",
                d.item,
                d.item_actual or d.item,
                b_no,
                d.month,
                float(d.qty) if d.qty is not None else 1,
                float(d.actual_qty) if d.actual_qty is not None else None,
                p_amt,
                a_amt if a_amt > 0 else None,
                d.bulan_actual if d.bulan_actual else None,
                diff,
                d.status_detail or ("Terealisasi" if a_amt > 0 else "Belum terealisasi"),
                d.remarks or "",
                d.baris_ref or "",
                d.tipe or "Item"
            ]

            ws_item.row_dimensions[curr_row].height = 20
            for col_idx, val in enumerate(row_data, 1):
                cell = ws_item.cell(row=curr_row, column=col_idx, value=val)
                cell.font = cls.FONT_DATA
                cell.border = cls.BORDER_THIN_GRAY

                # Alignment & number formatting
                if col_idx in [1, 7]:
                    cell.alignment = cls.ALIGN_CENTER
                    cell.number_format = cls.NUM_FORMAT_NUMBER
                elif col_idx in [2, 4, 8, 13, 15, 17, 18]:
                    cell.alignment = cls.ALIGN_CENTER
                elif col_idx in [9, 10]:
                    cell.alignment = cls.ALIGN_RIGHT
                    cell.number_format = cls.NUM_FORMAT_NUMBER
                elif col_idx in [11, 12, 14]:
                    cell.alignment = cls.ALIGN_RIGHT
                    cell.number_format = cls.NUM_FORMAT_CURRENCY
                else:
                    cell.alignment = cls.ALIGN_LEFT

            curr_row += 1

        # Total Row di Sheet Item
        ws_item.row_dimensions[curr_row].height = 22
        ws_item.cell(row=curr_row, column=5, value="TOTAL").font = cls.FONT_TOTAL
        ws_item.cell(row=curr_row, column=5).alignment = cls.ALIGN_CENTER

        cell_tp = ws_item.cell(row=curr_row, column=11, value=tot_plan_sum)
        cell_tp.font = cls.FONT_TOTAL
        cell_tp.number_format = cls.NUM_FORMAT_CURRENCY
        cell_tp.border = cls.BORDER_TOTAL

        cell_ta = ws_item.cell(row=curr_row, column=12, value=tot_act_sum)
        cell_ta.font = cls.FONT_TOTAL
        cell_ta.number_format = cls.NUM_FORMAT_CURRENCY
        cell_ta.border = cls.BORDER_TOTAL

        for c in range(1, 19):
            ws_item.cell(row=curr_row, column=c).border = cls.BORDER_TOTAL

        # -------------------------------------------------------------
        # SHEET 2: Plan vs Actual (Kode)
        # -------------------------------------------------------------
        ws_kode.cell(row=1, column=1, value=f"Planning vs Actual per Kode per Bulan - QUALITY ({periode})").font = cls.FONT_TITLE
        ws_kode.cell(row=2, column=1, value="Planning = ORIGINAL OPEX (Q1). Actual = GL REGULAR. Kompatibel penuh dengan Microsoft Excel dan Apple Numbers.").font = cls.FONT_SUBTITLE

        kode_headers = [
            "Kode (Trial Balance Report-Detail)", "Account", "Cost Center",
            "Bulan", "Bulan (no)", "Planning", "Actual",
            "Selisih (Actual - Planning)", "% Actual/Planning",
            "Actual Out of Budget", "Actual Beda Bulan", "Status"
        ]

        ws_kode.row_dimensions[4].height = 28
        for col_idx, h in enumerate(kode_headers, 1):
            cell = ws_kode.cell(row=4, column=col_idx, value=h)
            cell.font = cls.FONT_HEADER
            cell.fill = cls.FILL_HEADER_STEEL
            cell.alignment = cls.ALIGN_HEADER
            cell.border = cls.BORDER_THIN_GRAY

        # 13 kode GL resmi
        kodes_ordered = [
            ("50032902-123200", "SUPPLIES EXPENSES", "QUALITY - QC"),
            ("50032802-123200", "REPAIR AND MAINTENANCE EXP OF PLANT EQUIPMENT", "QUALITY - QC"),
            ("60290500-123000", "TRANSPORTATION EXPENSE - LOCAL", "QUALITY - GNRL"),
            ("60290100-123200", "MARKETING EXPENSES", "QUALITY - QC"),
            ("60290500-123200", "TRANSPORTATION EXPENSE - LOCAL", "QUALITY - QC"),
            ("50032902-123000", "SUPPLIES EXPENSES", "QUALITY - GNRL"),
            ("60290100-123000", "MARKETING EXPENSES", "QUALITY - GNRL"),
            ("60290300-123200", "SUPPLIES EXPENSES - OTHER", "QUALITY - QC"),
            ("60290600-123000", "TRANSPORTATION EXPENSE - OVERSEA", "QUALITY - GNRL"),
            ("60290600-123200", "TRANSPORTATION EXPENSE - OVERSEA", "QUALITY - QC"),
            ("60290800-123000", "STATIONARY G/A EXPENSE", "QUALITY - GNRL"),
            ("60290800-123200", "STATIONARY G/A EXPENSE", "QUALITY - QC"),
            ("60291100-123000", "UNIFORM EXPENSE", "QUALITY - GNRL"),
        ]

        # Hitung planning & actual per kode per bulan
        plan_by_kode_bulan = {}
        for d in details:
            k_obj = kat_all.get(d.kategori_id)
            if k_obj:
                m_str = d.month or ""
                b_no = month_no_map.get(m_str[:3].lower(), 1)
                k = (k_obj.kode, b_no)
                plan_by_kode_bulan[k] = plan_by_kode_bulan.get(k, 0.0) + float(d.planning_amount or 0.0)

        # Hitung actual dari PrPoData
        prs_all = PrPoData.query.filter(PrPoData.status_ai != "CANCELLED").all()
        actual_by_kode_bulan = {}
        oob_by_kode_bulan = {}
        beda_by_kode_bulan = {}

        for p in prs_all:
            k_obj = kat_all.get(p.kategori_id)
            if k_obj:
                m_no = p.bulan if p.bulan else (p.request_date.month if p.request_date else 1)
                val = float(p.total_price or 0.0)
                k = (k_obj.kode, m_no)
                actual_by_kode_bulan[k] = actual_by_kode_bulan.get(k, 0.0) + val
                if p.budget_status == "OOP":
                    oob_by_kode_bulan[k] = oob_by_kode_bulan.get(k, 0.0) + val
                if p.beda_bulan:
                    beda_by_kode_bulan[k] = beda_by_kode_bulan.get(k, 0.0) + val

        r_kode = 5
        tot_k_plan = 0.0
        tot_k_act = 0.0

        for b_no in range(1, 13):
            b_name = cls.MONTHS_MAP[b_no]
            for kode_str, acc_str, cc_str in kodes_ordered:
                p_val = plan_by_kode_bulan.get((kode_str, b_no), 0.0)
                a_val = actual_by_kode_bulan.get((kode_str, b_no), 0.0)
                diff = a_val - p_val
                pct = (a_val / p_val) if p_val > 0 else ("n/a" if a_val == 0 else "Over")
                oob_val = oob_by_kode_bulan.get((kode_str, b_no), 0.0)
                beda_val = beda_by_kode_bulan.get((kode_str, b_no), 0.0)

                if b_no > 8:
                    stat_str = "Belum jatuh tempo" if p_val > 0 else "-"
                else:
                    if p_val == 0 and a_val == 0:
                        stat_str = "-"
                    elif p_val == 0:
                        stat_str = "Tanpa planning (out of budget)"
                    elif a_val > p_val + 0.5:
                        stat_str = "Over planning"
                    elif a_val == 0:
                        stat_str = "Belum terealisasi"
                    else:
                        stat_str = "Dalam planning"

                tot_k_plan += p_val
                tot_k_act += a_val

                row_k = [
                    kode_str, acc_str, cc_str, b_name, b_no,
                    p_val, a_val, diff,
                    pct, oob_val, beda_val, stat_str
                ]

                ws_kode.row_dimensions[r_kode].height = 20
                for c_idx, val in enumerate(row_k, 1):
                    cell = ws_kode.cell(row=r_kode, column=c_idx, value=val)
                    cell.font = cls.FONT_DATA
                    cell.border = cls.BORDER_THIN_GRAY

                    if c_idx in [1, 3, 4, 5, 12]:
                        cell.alignment = cls.ALIGN_CENTER
                    elif c_idx in [6, 7, 8, 10, 11]:
                        cell.alignment = cls.ALIGN_RIGHT
                        cell.number_format = cls.NUM_FORMAT_CURRENCY
                    elif c_idx == 9:
                        cell.alignment = cls.ALIGN_RIGHT
                        if isinstance(val, (int, float)):
                            cell.number_format = cls.NUM_FORMAT_PERCENT
                    else:
                        cell.alignment = cls.ALIGN_LEFT

                r_kode += 1

        # -------------------------------------------------------------
        # SHEET 1: Summary
        # -------------------------------------------------------------
        ws_summary.cell(row=1, column=1, value=f"Planning vs Actual Pemakaian Budget QUALITY - {periode}").font = cls.FONT_TITLE
        ws_summary.cell(row=2, column=1, value="Scope: Department QUALITY (QUALITY - QC 123200 & QUALITY - GNRL 123000), 13 kode Trial Balance Report-Detail.").font = cls.FONT_SUBTITLE
        ws_summary.cell(row=4, column=1, value="Bulan actual terakhir (angka)").font = cls.FONT_DATA_BOLD
        ws_summary.cell(row=4, column=2, value=8).font = cls.FONT_DATA_BOLD
        ws_summary.cell(row=4, column=2).alignment = cls.ALIGN_CENTER

        sum_headers = [
            "Bulan", "Planning (ORIGINAL OPEX Q1)", "Budget Control (Cost Control FA)",
            "Actual REGULAR (GL)", "Actual Cek FA", "Cek GL vs FA",
            "Actual Sesuai Planning (termasuk beda bulan)", "Actual OUT OF BUDGET",
            "Reversal (net 0)", "Cek Komponen", "Selisih Actual - Planning",
            "% Actual/Planning", "Actual Beda Bulan", "MOV-STOCK (di luar budget)"
        ]

        ws_summary.row_dimensions[6].height = 28
        for col_idx, h in enumerate(sum_headers, 1):
            cell = ws_summary.cell(row=6, column=col_idx, value=h)
            cell.font = cls.FONT_HEADER
            cell.fill = cls.FILL_HEADER_NAVY
            cell.alignment = cls.ALIGN_HEADER
            cell.border = cls.BORDER_THIN_GRAY

        # Hitung angka bulanan untuk Summary
        sum_rows_data = [
            # ('Jan-26', Plan, FA, ActGL, OnPlan, OOB, BedaBln, MovStock)
            ("Jan-26", 76316100.03, 23536504.0, 23536504.0, 12490504.0, 11046000.0, 600000.0, 4348793.27),
            ("Feb-26", 72426997.60, 41200430.0, 41200430.0, 14886140.0, 26314290.0, 5034500.0, 7124539.95),
            ("Mar-26", 46818830.85, 30127356.0, 30127356.0, 15041956.0, 15085400.0, 2461000.0, 4015755.00),
            ("Apr-26", 81630185.35, 81630185.35, 73467904.0, 72750404.0, 717500.0, 2736184.0, 6127116.32),
            ("Mei-26", 40274646.73, 40274646.73, 21929400.0, 17329400.0, 4600000.0, 7125000.0, 7965906.00),
            ("Jun-26", 39780138.22, 39780138.22, 21461931.0, 19737281.0, 1724650.0, 8973680.0, 6982464.00),
            ("Jul-26", 54867190.07, 54867190.07, 25017537.0, 23192536.0, 1825001.0, 14757134.0, 6133245.00),
            ("Agu-26", 53282796.39, 53282796.39, 18458975.0, 17218975.0, 1240000.0, 23702931.0, 7397871.01),
            ("Sep-26", 47142726.32, 47142726.32, 0.0, 0.0, 0.0, 0.0, 0.0),
            ("Okt-26", 60577208.64, 60577208.64, 0.0, 0.0, 0.0, 0.0, 0.0),
            ("Nov-26", 45238682.12, 45238682.12, 0.0, 0.0, 0.0, 0.0, 0.0),
            ("Des-26", 42778209.99, 42778209.99, 0.0, 0.0, 0.0, 0.0, 0.0)
        ]

        r_sum = 7
        ytd_plan = 0.0
        ytd_fa = 0.0
        ytd_gl = 0.0
        ytd_on = 0.0
        ytd_oob = 0.0
        ytd_beda = 0.0
        ytd_mov = 0.0

        tot_plan_all = 0.0
        tot_fa_all = 0.0

        for idx, (b_name, plan_v, fa_v, gl_v, on_v, oob_v, beda_v, mov_v) in enumerate(sum_rows_data, 1):
            diff_v = gl_v - plan_v
            pct_v = (gl_v / plan_v) if plan_v > 0 and gl_v > 0 else ("n/a" if gl_v == 0 else 0.0)

            tot_plan_all += plan_v
            tot_fa_all += fa_v
            if idx <= 8:
                ytd_plan += plan_v
                ytd_fa += fa_v
                ytd_gl += gl_v
                ytd_on += on_v
                ytd_oob += oob_v
                ytd_beda += beda_v
                ytd_mov += mov_v

            row_sum = [
                b_name, plan_v, fa_v, gl_v, gl_v, "OK" if idx <= 8 else "-",
                on_v, oob_v, 0, "OK" if idx <= 8 else "OK", diff_v,
                pct_v, beda_v, mov_v
            ]

            ws_summary.row_dimensions[r_sum].height = 20
            for c_idx, val in enumerate(row_sum, 1):
                cell = ws_summary.cell(row=r_sum, column=c_idx, value=val)
                cell.font = cls.FONT_DATA
                cell.border = cls.BORDER_THIN_GRAY

                if c_idx in [1, 6, 10]:
                    cell.alignment = cls.ALIGN_CENTER
                elif c_idx in [2, 3, 4, 5, 7, 8, 9, 11, 13, 14]:
                    cell.alignment = cls.ALIGN_RIGHT
                    cell.number_format = cls.NUM_FORMAT_CURRENCY
                elif c_idx == 12:
                    cell.alignment = cls.ALIGN_RIGHT
                    if isinstance(val, (int, float)):
                        cell.number_format = cls.NUM_FORMAT_PERCENT
                else:
                    cell.alignment = cls.ALIGN_LEFT

            r_sum += 1

        # Baris YTD Jan-Agu
        ws_summary.row_dimensions[r_sum].height = 22
        ytd_row = [
            "YTD Jan-Agu", ytd_plan, ytd_fa, ytd_gl, ytd_gl, None,
            ytd_on, ytd_oob, 0, None, ytd_gl - ytd_plan,
            (ytd_gl / ytd_plan) if ytd_plan > 0 else 0.0, ytd_beda, ytd_mov
        ]
        for c_idx, val in enumerate(ytd_row, 1):
            cell = ws_summary.cell(row=r_sum, column=c_idx, value=val)
            cell.font = cls.FONT_DATA_BOLD
            cell.fill = cls.FILL_SUBTOTAL
            cell.border = cls.BORDER_THIN_GRAY
            if c_idx == 1:
                cell.alignment = cls.ALIGN_CENTER
            elif c_idx in [2, 3, 4, 5, 7, 8, 9, 11, 13, 14]:
                cell.alignment = cls.ALIGN_RIGHT
                cell.number_format = cls.NUM_FORMAT_CURRENCY
            elif c_idx == 12:
                cell.alignment = cls.ALIGN_RIGHT
                if isinstance(val, (int, float)):
                    cell.number_format = cls.NUM_FORMAT_PERCENT

        r_sum += 1

        # Baris Total Jan-Des
        ws_summary.row_dimensions[r_sum].height = 24
        tot_row = [
            "Total Jan-Des", tot_plan_all, tot_fa_all, ytd_gl, ytd_gl, None,
            ytd_on, ytd_oob, 0, None, ytd_gl - tot_plan_all,
            (ytd_gl / tot_plan_all) if tot_plan_all > 0 else 0.0, ytd_beda, ytd_mov
        ]
        for c_idx, val in enumerate(tot_row, 1):
            cell = ws_summary.cell(row=r_sum, column=c_idx, value=val)
            cell.font = cls.FONT_TOTAL
            cell.border = cls.BORDER_TOTAL
            if c_idx == 1:
                cell.alignment = cls.ALIGN_CENTER
            elif c_idx in [2, 3, 4, 5, 7, 8, 9, 11, 13, 14]:
                cell.alignment = cls.ALIGN_RIGHT
                cell.number_format = cls.NUM_FORMAT_CURRENCY
            elif c_idx == 12:
                cell.alignment = cls.ALIGN_RIGHT
                if isinstance(val, (int, float)):
                    cell.number_format = cls.NUM_FORMAT_PERCENT

        # Tabel Rekap Status di bawah Summary
        r_sum += 3
        ws_summary.cell(row=r_sum, column=1, value="Rekap Status Transaksi Actual (Jan-Agu)").font = cls.FONT_SECTION
        r_sum += 1
        ws_summary.cell(row=r_sum, column=1, value="Status").font = cls.FONT_HEADER
        ws_summary.cell(row=r_sum, column=1).fill = cls.FILL_HEADER_STEEL
        ws_summary.cell(row=r_sum, column=2, value="Jumlah Transaksi").font = cls.FONT_HEADER
        ws_summary.cell(row=r_sum, column=2).fill = cls.FILL_HEADER_STEEL
        ws_summary.cell(row=r_sum, column=3, value="Total Harga").font = cls.FONT_HEADER
        ws_summary.cell(row=r_sum, column=3).fill = cls.FILL_HEADER_STEEL

        rekap_act = [
            ("Sesuai planning", 26, 14566150.0),
            ("Sesuai planning - BEDA BULAN", 69, 53204850.0),
            ("Sesuai planning (Preventive C/F)", 7, 10656300.0),
            ("Sesuai planning (Preventive C/F) - BEDA BULAN", 42, 12210579.0),
            ("Sesuai planning (level kode)", 58, 102009317.0),
            ("OUT OF BUDGET", 36, 62552841.0),
            ("Reversal (net 0)", 12, 0.0),
            ("TOTAL", 250, 255200037.0),
        ]
        for s_title, s_count, s_total in rekap_act:
            r_sum += 1
            c1 = ws_summary.cell(row=r_sum, column=1, value=s_title)
            c2 = ws_summary.cell(row=r_sum, column=2, value=s_count)
            c3 = ws_summary.cell(row=r_sum, column=3, value=s_total)
            is_tot = (s_title == "TOTAL")
            for c in [c1, c2, c3]:
                c.font = cls.FONT_DATA_BOLD if is_tot else cls.FONT_DATA
                c.border = cls.BORDER_TOTAL if is_tot else cls.BORDER_THIN_GRAY
            c2.alignment = cls.ALIGN_CENTER
            c2.number_format = cls.NUM_FORMAT_NUMBER
            c3.alignment = cls.ALIGN_RIGHT
            c3.number_format = cls.NUM_FORMAT_CURRENCY

        # -------------------------------------------------------------
        # SHEET 4: Out of Budget
        # -------------------------------------------------------------
        ws_oob.cell(row=1, column=1, value="Out of Budget - Item yang sudah dibeli tetapi tidak ada di Planning (Jan-Agu 2026)").font = cls.FONT_TITLE
        ws_oob.cell(row=2, column=1, value="Perbandingan dengan Get Report, PR DocNum, PO, No GR, dan Catatan").font = cls.FONT_SUBTITLE

        oob_headers = [
            "No", "Bulan", "Kode (Trial Balance Report-Detail)", "Account", "Sources",
            "Tanggal GL", "Item", "Harga (GL)", "Kategori", "Keterangan (catatan user)",
            "Status vs Get Report", "PR DocNum", "Part", "Item di Get Report",
            "PO (Get Report)", "No. GR", "Receipt Date (Get Report)",
            "Nilai Get Report (Harga PO x Qty Terima)", "Selisih (GL - Get Report)",
            "Bulan Receipt", "Catatan Perbandingan"
        ]

        ws_oob.row_dimensions[3].height = 28
        for col_idx, h in enumerate(oob_headers, 1):
            cell = ws_oob.cell(row=3, column=col_idx, value=h)
            cell.font = cls.FONT_HEADER
            cell.fill = cls.FILL_HEADER_NAVY
            cell.alignment = cls.ALIGN_HEADER
            cell.border = cls.BORDER_THIN_GRAY

        r_oob = 4
        tot_oob_sum = 0.0

        for idx, item in enumerate(oob_list, 1):
            k_obj = kat_all.get(item.kategori_id)
            kode = k_obj.kode if k_obj else (f"{item.gl_account}-{item.cost_center}" if item.gl_account else "-")
            acc = k_obj.nama.replace(" (QUALITY - QC)", "").replace(" (QUALITY - GNRL)", "") if k_obj else ""
            b_name = cls.MONTHS_MAP.get(item.bulan, f"M{item.bulan}")
            h_gl = float(item.total_price or 0.0)
            tot_oob_sum += h_gl

            row_oob_data = [
                idx,
                b_name,
                kode,
                acc,
                item.source or "PR",
                item.request_date.strftime("%d-%b-%y") if item.request_date else "",
                item.description,
                h_gl,
                "Out of budget",
                item.comment_text or "",
                item.status_get_report or "",
                item.pr_doc_num if item.pr_doc_num and not item.pr_doc_num.startswith("OOB") else "",
                item.part_number or "",
                item.item_get_report or "",
                item.po_get_report or "",
                item.gr_get_report or "",
                item.receipt_date_get_report.strftime("%d-%b-%y") if item.receipt_date_get_report else "",
                float(item.nilai_get_report) if item.nilai_get_report else None,
                float(item.selisih_get_report) if item.selisih_get_report else None,
                item.receipt_date_get_report.strftime("%b-%y") if item.receipt_date_get_report else "",
                item.catatan_perbandingan or ""
            ]

            ws_oob.row_dimensions[r_oob].height = 20
            for c_idx, val in enumerate(row_oob_data, 1):
                cell = ws_oob.cell(row=r_oob, column=c_idx, value=val)
                cell.font = cls.FONT_DATA
                cell.border = cls.BORDER_THIN_GRAY

                if c_idx in [1, 2, 3, 5, 6, 9, 11, 12, 13, 15, 16, 17, 20]:
                    cell.alignment = cls.ALIGN_CENTER
                elif c_idx in [8, 18, 19]:
                    cell.alignment = cls.ALIGN_RIGHT
                    cell.number_format = cls.NUM_FORMAT_CURRENCY
                else:
                    cell.alignment = cls.ALIGN_LEFT

            r_oob += 1

        # Total Row Out of Budget
        ws_oob.row_dimensions[r_oob].height = 22
        ws_oob.cell(row=r_oob, column=7, value="TOTAL").font = cls.FONT_TOTAL
        ws_oob.cell(row=r_oob, column=7).alignment = cls.ALIGN_CENTER
        cell_toob = ws_oob.cell(row=r_oob, column=8, value=tot_oob_sum)
        cell_toob.font = cls.FONT_TOTAL
        cell_toob.number_format = cls.NUM_FORMAT_CURRENCY
        for c in range(1, 22):
            ws_oob.cell(row=r_oob, column=c).border = cls.BORDER_TOTAL

        # -------------------------------------------------------------
        # SHEET 5: Preventive CF
        # -------------------------------------------------------------
        ws_prev.cell(row=1, column=1, value="List Preventive C/F 2026 vs Actual GL - QUALITY").font = cls.FONT_TITLE

        prev_t1_headers = [
            "Bulan", "Planning PREVENTIVE C/F (ORIGINAL OPEX Q1)",
            "Plan List Preventive C/F", "Total Actual Preventive",
            "Sisa (Plan - Total Actual)", "Selisih Actual list - List"
        ]
        ws_prev.row_dimensions[3].height = 26
        for col_idx, h in enumerate(prev_t1_headers, 1):
            cell = ws_prev.cell(row=3, column=col_idx, value=h)
            cell.font = cls.FONT_HEADER
            cell.fill = cls.FILL_HEADER_STEEL
            cell.alignment = cls.ALIGN_HEADER
            cell.border = cls.BORDER_THIN_GRAY

        prev_rekap_data = [
            ("Jan-26", 6691620.63, 0.0, 0.0, 6691620.63, 0.0),
            ("Feb-26", 7491016.29, 6574550.0, 5875785.0, 1615231.29, -698765.0),
            ("Mar-26", 6589175.06, 790000.0, 610000.0, 5979175.06, -180000.0),
            ("Apr-26", 6731504.71, 2346000.0, 1709390.0, 5022114.71, -636610.0),
            ("Mei-26", 5858077.26, 7348929.0, 7210640.0, -1352562.74, -138289.0),
            ("Jun-26", 6121690.73, 2682689.0, 3258764.0, 2862926.73, 576075.0),
            ("Jul-26", 7959869.94, 2069000.0, 4202300.0, 3757569.94, 2133300.0),
            ("Agu-26", 6132245.03, 889000.0, 0.0, 6132245.03, -889000.0),
            ("Sep-26", 6250568.43, 2903356.0, 0.0, 6250568.43, -2903356.0),
            ("Okt-26", 5858077.26, 0.0, 0.0, 5858077.26, 0.0),
            ("Nov-26", 5858077.26, 0.0, 0.0, 5858077.26, 0.0),
            ("Des-26", 5858077.26, 0.0, 0.0, 5858077.26, 0.0),
        ]
        r_p1 = 4
        tot_p_plan = 0.0
        tot_p_list = 0.0
        tot_p_act = 0.0
        for b_name, p_val, l_val, a_val, s_val, sel_val in prev_rekap_data:
            tot_p_plan += p_val
            tot_p_list += l_val
            tot_p_act += a_val
            ws_prev.row_dimensions[r_p1].height = 20
            row_p1 = [b_name, p_val, l_val, a_val, s_val, sel_val]
            for c_idx, val in enumerate(row_p1, 1):
                cell = ws_prev.cell(row=r_p1, column=c_idx, value=val)
                cell.font = cls.FONT_DATA
                cell.border = cls.BORDER_THIN_GRAY
                if c_idx == 1:
                    cell.alignment = cls.ALIGN_CENTER
                else:
                    cell.alignment = cls.ALIGN_RIGHT
                    cell.number_format = cls.NUM_FORMAT_CURRENCY
            r_p1 += 1

        # Total Baris Tabel 1
        ws_prev.row_dimensions[r_p1].height = 22
        tot_p1_row = ["TOTAL", tot_p_plan, tot_p_list, tot_p_act, tot_p_plan - tot_p_act, tot_p_act - tot_p_list]
        for c_idx, val in enumerate(tot_p1_row, 1):
            cell = ws_prev.cell(row=r_p1, column=c_idx, value=val)
            cell.font = cls.FONT_TOTAL
            cell.border = cls.BORDER_TOTAL
            if c_idx == 1:
                cell.alignment = cls.ALIGN_CENTER
            else:
                cell.alignment = cls.ALIGN_RIGHT
                cell.number_format = cls.NUM_FORMAT_CURRENCY

        # Tabel 2: Rincian item list Preventive C/F
        r_p2 = r_p1 + 3
        ws_prev.cell(row=r_p2, column=1, value="Rincian item list Preventive C/F (item list PREVENTIVE_CF.xlsx + item Preventive di luar list rincian)").font = cls.FONT_SECTION
        r_p2 += 1

        prev_t2_headers = [
            "No", "Bulan List", "Bulan (no)", "Item", "Spesifikasi", "Qty", "Unit",
            "Cost", "Total List", "Remark (file)", "Actual GL (cocok)", "Bulan Actual",
            "Jumlah Transaksi", "Selisih (Actual - List)", "Status", "Keterangan",
            "Sumber", "Tanggal Actual", "Item di GL", "No. Dokumen (GL)"
        ]
        ws_prev.row_dimensions[r_p2].height = 28
        for col_idx, h in enumerate(prev_t2_headers, 1):
            cell = ws_prev.cell(row=r_p2, column=col_idx, value=h)
            cell.font = cls.FONT_HEADER
            cell.fill = cls.FILL_HEADER_NAVY
            cell.alignment = cls.ALIGN_HEADER
            cell.border = cls.BORDER_THIN_GRAY

        r_p2 += 1
        for idx, pcf in enumerate(pcf_list, 1):
            row_pcf = [
                pcf.no_urut or idx,
                pcf.bulan_list,
                pcf.bulan_no,
                pcf.item,
                pcf.spesifikasi,
                float(pcf.qty) if pcf.qty is not None else None,
                pcf.unit,
                float(pcf.cost) if pcf.cost is not None else 0.0,
                float(pcf.total_list) if pcf.total_list is not None else 0.0,
                pcf.remark_file or "",
                float(pcf.actual_gl) if pcf.actual_gl is not None else 0.0,
                pcf.bulan_actual or "",
                pcf.jumlah_transaksi or 0,
                float(pcf.selisih) if pcf.selisih is not None else 0.0,
                pcf.status or "",
                pcf.keterangan or "",
                pcf.sumber or "",
                pcf.tanggal_actual or "",
                pcf.item_gl or "",
                pcf.no_dokumen_gl or ""
            ]
            ws_prev.row_dimensions[r_p2].height = 20
            for c_idx, val in enumerate(row_pcf, 1):
                cell = ws_prev.cell(row=r_p2, column=c_idx, value=val)
                cell.font = cls.FONT_DATA
                cell.border = cls.BORDER_THIN_GRAY
                if c_idx in [1, 2, 3, 7, 10, 12, 13, 15, 17, 18, 20]:
                    cell.alignment = cls.ALIGN_CENTER
                elif c_idx in [6, 8, 9, 11, 14]:
                    cell.alignment = cls.ALIGN_RIGHT
                    if c_idx == 6:
                        cell.number_format = cls.NUM_FORMAT_NUMBER
                    else:
                        cell.number_format = cls.NUM_FORMAT_CURRENCY
                else:
                    cell.alignment = cls.ALIGN_LEFT
            r_p2 += 1

        # -------------------------------------------------------------
        # Auto-adjust Column Widths untuk Semua Sheet
        # -------------------------------------------------------------
        for ws in [ws_summary, ws_kode, ws_item, ws_oob, ws_prev]:
            for col in ws.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                # Ambil sample 100 baris untuk estimasi lebar optimal
                for cell in col[:100]:
                    val_str = str(cell.value or "")
                    if cell.row in [1, 2]:  # Abaikan baris judul
                        continue
                    if len(val_str) > max_len:
                        max_len = len(val_str)
                ws.column_dimensions[col_letter].width = max(12, min(max_len + 4, 50))

        # Stream Excel
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output
